import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------
# Presentation Setup (16:9 Widescreen)
# ----------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]  # Blank layout

# Color Palette
DARK_BG = RGBColor(15, 23, 42)      # Slate 900
LIGHT_BG = RGBColor(248, 250, 252) # Slate 50
CARD_BG = RGBColor(255, 255, 255)  # Pure White
CARD_BORDER = RGBColor(226, 232, 240) # Slate 200

EMERALD_DARK = RGBColor(4, 120, 87)  # Emerald 700
EMERALD_PRIMARY = RGBColor(5, 150, 105) # Emerald 600
EMERALD_LIGHT = RGBColor(209, 250, 229) # Emerald 100
EMERALD_TEXT = RGBColor(6, 78, 59)   # Emerald 900

AMBER_PRIMARY = RGBColor(217, 119, 6) # Amber 600
AMBER_LIGHT = RGBColor(254, 243, 199) # Amber 100
AMBER_TEXT = RGBColor(120, 53, 15)   # Amber 900

SLATE_DARK = RGBColor(30, 41, 59)   # Slate 800
SLATE_MUTED = RGBColor(100, 116, 139) # Slate 500

# Screenshot Paths
ARTIFACT_DIR = r"C:\Users\MINHA FATHIMA C\.gemini\antigravity-ide\brain\9df0dcab-291d-4dfd-b2d9-9a19395e2bab"
HOME_IMG = os.path.join(ARTIFACT_DIR, "home_page_1789806908893.png")
SDG_IMG = os.path.join(ARTIFACT_DIR, "sdg_impact_page_1789806964309.png")
ANALYZE_IMG = os.path.join(ARTIFACT_DIR, "analyze_page_1789806987399.png")
HISTORY_IMG = os.path.join(ARTIFACT_DIR, "history_page_1789807011469.png")
DASHBOARD_IMG = os.path.join(ARTIFACT_DIR, "dashboard_page_1789807034513.png")

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, tag_text, title_text, category_badge=None):
    # Header container
    top = Inches(0.4)
    
    # Tag / Category Badge
    badge_box = slide.shapes.add_textbox(Inches(0.8), top, Inches(11.733), Inches(0.35))
    tf_b = badge_box.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = tag_text.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = EMERALD_PRIMARY
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), top + Inches(0.3), Inches(11.733), Inches(0.6))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = SLATE_DARK

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

# ====================================================
# SLIDE 1: Title Slide (Dark Theme)
# ====================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1, DARK_BG)

# Title & Subtitle Card
t_card = create_card(slide1, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5), bg_color=RGBColor(30, 41, 59), border_color=EMERALD_PRIMARY)

tf1 = t_card.text_frame
tf1.word_wrap = True
tf1.margin_left = Inches(0.8)
tf1.margin_top = Inches(0.8)
tf1.margin_right = Inches(0.8)

p1 = tf1.paragraphs[0]
p1.text = "SecondLife AI"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = RGBColor(255, 255, 255)

p2 = tf1.add_paragraph()
p2.text = "AI-Powered Sustainable Asset Recycling & Waste Reduction Platform"
p2.font.size = Pt(20)
p2.font.color.rgb = EMERALD_PRIMARY
p2.space_before = Pt(12)

p3 = tf1.add_paragraph()
p3.text = "Automated item classification, material circularity scoring, condition assessment, and community sharing to support global UN Sustainable Development Goals."
p3.font.size = Pt(14)
p3.font.color.rgb = RGBColor(203, 213, 225)
p3.space_before = Pt(20)

# Footer Note
f_box = slide1.shapes.add_textbox(Inches(1.5), Inches(6.3), Inches(10.333), Inches(0.5))
p_f = f_box.text_frame.paragraphs[0]
p_f.text = "Presented by SecondLife AI Project Team • Powered by React, FastAPI & PyTorch Vision"
p_f.font.size = Pt(11)
p_f.font.color.rgb = SLATE_MUTED
p_f.alignment = PP_ALIGN.CENTER

# ====================================================
# SLIDE 2: Project Name & Aligned UN SDGs (Mandatory Required Slide)
# ====================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2, LIGHT_BG)
add_header(slide2, "Core Purpose & Global Framework", "Project Name & Aligned UN Sustainable Development Goals (SDGs)")

# Left Column: Project Identity Card
create_card(slide2, Inches(0.8), Inches(1.4), Inches(3.8), Inches(5.5), bg_color=CARD_BG)
id_box = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(3.4), Inches(5.1))
tf_id = id_box.text_frame
tf_id.word_wrap = True

p_n1 = tf_id.paragraphs[0]
p_n1.text = "PROJECT NAME"
p_n1.font.size = Pt(11)
p_n1.font.bold = True
p_n1.font.color.rgb = EMERALD_PRIMARY

p_n2 = tf_id.add_paragraph()
p_n2.text = "SecondLife AI"
p_n2.font.size = Pt(28)
p_n2.font.bold = True
p_n2.font.color.rgb = SLATE_DARK
p_n2.space_before = Pt(4)

p_n3 = tf_id.add_paragraph()
p_n3.text = "Vision Statement"
p_n3.font.size = Pt(13)
p_n3.font.bold = True
p_n3.font.color.rgb = SLATE_DARK
p_n3.space_before = Pt(16)

p_n4 = tf_id.add_paragraph()
p_n4.text = "To eliminate unnecessary landfill disposal by equipping individuals and communities with instant artificial intelligence tools to classify, repair, reuse, or recycle household items."
p_n4.font.size = Pt(11)
p_n4.font.color.rgb = SLATE_MUTED
p_n4.space_before = Pt(6)

p_n5 = tf_id.add_paragraph()
p_n5.text = "Target Impact"
p_n5.font.size = Pt(13)
p_n5.font.bold = True
p_n5.font.color.rgb = SLATE_DARK
p_n5.space_before = Pt(16)

p_n6 = tf_id.add_paragraph()
p_n6.text = "• Substantial waste diversion\n• Lifetime product extension\n• Local community sharing\n• Closed-loop circular economy"
p_n6.font.size = Pt(11)
p_n6.font.color.rgb = SLATE_DARK
p_n6.space_before = Pt(6)

# Right Column: 3 SDG Cards
# SDG 12 Card (Primary)
c_sdg12 = create_card(slide2, Inches(4.8), Inches(1.4), Inches(7.733), Inches(1.7), bg_color=AMBER_LIGHT, border_color=AMBER_PRIMARY)
t_sdg12 = slide2.shapes.add_textbox(Inches(5.0), Inches(1.5), Inches(7.333), Inches(1.5))
tf_12 = t_sdg12.text_frame
tf_12.word_wrap = True

p12_1 = tf_12.paragraphs[0]
p12_1.text = "PRIMARY FOCUS • UN SDG 12: RESPONSIBLE CONSUMPTION & PRODUCTION"
p12_1.font.size = Pt(11)
p12_1.font.bold = True
p12_1.font.color.rgb = AMBER_TEXT

p12_2 = tf_12.add_paragraph()
p12_2.text = "Target 12.5: Substantially reduce waste generation through prevention, reduction, recycling, and reuse."
p12_2.font.size = Pt(13)
p12_2.font.bold = True
p12_2.font.color.rgb = SLATE_DARK
p12_2.space_before = Pt(4)

p12_3 = tf_12.add_paragraph()
p12_3.text = "SecondLife AI gives users immediate, actionable guidance on what to do with unwanted items instead of throwing them into landfills, maximizing material circularity."
p12_3.font.size = Pt(11)
p12_3.font.color.rgb = SLATE_DARK
p12_3.space_before = Pt(4)

# SDG 11 Card (Supporting)
c_sdg11 = create_card(slide2, Inches(4.8), Inches(3.3), Inches(7.733), Inches(1.7), bg_color=CARD_BG)
t_sdg11 = slide2.shapes.add_textbox(Inches(5.0), Inches(3.4), Inches(7.333), Inches(1.5))
tf_11 = t_sdg11.text_frame
tf_11.word_wrap = True

p11_1 = tf_11.paragraphs[0]
p11_1.text = "SUPPORTING GOAL • UN SDG 11: SUSTAINABLE CITIES & COMMUNITIES"
p11_1.font.size = Pt(11)
p11_1.font.bold = True
p11_1.font.color.rgb = EMERALD_PRIMARY

p11_2 = tf_11.add_paragraph()
p11_2.text = "Target 11.6: Reduce the adverse per capita environmental impact of cities, including municipal waste management."
p11_2.font.size = Pt(13)
p11_2.font.bold = True
p11_2.font.color.rgb = SLATE_DARK
p11_2.space_before = Pt(4)

p11_3 = tf_11.add_paragraph()
p11_3.text = "Household waste segregation and local item sharing directly reduce municipal solid waste burdens and foster cleaner, more resilient urban communities."
p11_3.font.size = Pt(11)
p11_3.font.color.rgb = SLATE_MUTED
p11_3.space_before = Pt(4)

# SDG 13 Card (Supporting)
c_sdg13 = create_card(slide2, Inches(4.8), Inches(5.2), Inches(7.733), Inches(1.7), bg_color=CARD_BG)
t_sdg13 = slide2.shapes.add_textbox(Inches(5.0), Inches(5.3), Inches(7.333), Inches(1.5))
tf_13 = t_sdg13.text_frame
tf_13.word_wrap = True

p13_1 = tf_13.paragraphs[0]
p13_1.text = "SUPPORTING GOAL • UN SDG 13: CLIMATE ACTION"
p13_1.font.size = Pt(11)
p13_1.font.bold = True
p13_1.font.color.rgb = EMERALD_PRIMARY

p13_2 = tf_13.add_paragraph()
p13_2.text = "Target 13.3: Improve education and awareness on climate change mitigation."
p13_2.font.size = Pt(13)
p13_2.font.bold = True
p13_2.font.color.rgb = SLATE_DARK
p13_2.space_before = Pt(4)

p13_3 = tf_13.add_paragraph()
p13_3.text = "Extending item lifespans and recycling scrap materials avoids upstream industrial carbon emissions and prevents organic paper/cardboard landfill methane generation."
p13_3.font.size = Pt(11)
p13_3.font.color.rgb = SLATE_MUTED
p13_3.space_before = Pt(4)


# ====================================================
# SLIDE 3: Problem Statement & Solution Vision
# ====================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3, LIGHT_BG)
add_header(slide3, "Challenges & Innovative Solution", "Problem Statement & SecondLife AI Approach")

# Left Column: Problem Card
create_card(slide3, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=CARD_BG)
p_box = slide3.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.3), Inches(5.1))
tf_p = p_box.text_frame
tf_p.word_wrap = True

pp1 = tf_p.paragraphs[0]
pp1.text = "THE PROBLEM"
pp1.font.size = Pt(12)
pp1.font.bold = True
pp1.font.color.rgb = RGBColor(225, 29, 72) # Rose 600

pp2 = tf_p.add_paragraph()
pp2.text = "Linear 'Take-Make-Dispose' Economy"
pp2.font.size = Pt(20)
pp2.font.bold = True
pp2.font.color.rgb = SLATE_DARK
pp2.space_before = Pt(6)

pp3 = tf_p.add_paragraph()
pp3.text = "• Millions of usable household items are dumped into landfills annually due to lack of immediate disposal guidance."
pp3.font.size = Pt(13)
pp3.font.color.rgb = SLATE_DARK
pp3.space_before = Pt(14)

pp4 = tf_p.add_paragraph()
pp4.text = "• Consumers struggle to determine whether an item can be repaired, donated, upcycled, or recycled."
pp4.font.size = Pt(13)
pp4.font.color.rgb = SLATE_DARK
pp4.space_before = Pt(10)

pp5 = tf_p.add_paragraph()
pp5.text = "• Improper category classification leads to contaminated recycling streams and missed donation opportunities."
pp5.font.size = Pt(13)
pp5.font.color.rgb = SLATE_DARK
pp5.space_before = Pt(10)

# Right Column: Solution Card
create_card(slide3, Inches(6.833), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=EMERALD_LIGHT, border_color=EMERALD_PRIMARY)
s_box = slide3.shapes.add_textbox(Inches(7.033), Inches(1.6), Inches(5.3), Inches(5.1))
tf_s = s_box.text_frame
tf_s.word_wrap = True

sp1 = tf_s.paragraphs[0]
sp1.text = "OUR SOLUTION"
sp1.font.size = Pt(12)
sp1.font.bold = True
sp1.font.color.rgb = EMERALD_DARK

sp2 = tf_s.add_paragraph()
sp2.text = "SecondLife AI Circular Platform"
sp2.font.size = Pt(20)
sp2.font.bold = True
sp2.font.color.rgb = SLATE_DARK
sp2.space_before = Pt(6)

sp3 = tf_s.add_paragraph()
sp3.text = "1. AI Image & Text Classifier: Instantly normalizes items into 8 core categories (Electronics, Clothing, Paper, Plastic, Glass, Metal, Furniture, Other)."
sp3.font.size = Pt(13)
sp3.font.color.rgb = SLATE_DARK
sp3.space_before = Pt(14)

sp4 = tf_s.add_paragraph()
sp4.text = "2. Condition-Aware Recommendation Engine: Evaluates item state (Good, Fair, Damaged, Not usable) to suggest Repair, Reuse, Donation, or Scrap Recycling."
sp4.font.size = Pt(13)
sp4.font.color.rgb = SLATE_DARK
sp4.space_before = Pt(10)

sp5 = tf_s.add_paragraph()
sp5.text = "3. Digital Provenance & QR Tagging: Generates printable QR tags and item certificates for transparent community sharing."
sp5.font.size = Pt(13)
sp5.font.color.rgb = SLATE_DARK
sp5.space_before = Pt(10)


# ====================================================
# SLIDE 4: Home & Platform Navigation (With Screenshot)
# ====================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4, LIGHT_BG)
add_header(slide4, "User Interface & Experience", "Platform Landing & Navigation Overview")

# Left Column: Features List
create_card(slide4, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), bg_color=CARD_BG)
h_box = slide4.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
tf_h = h_box.text_frame
tf_h.word_wrap = True

hp1 = tf_h.paragraphs[0]
hp1.text = "KEY INTERFACE FEATURES"
hp1.font.size = Pt(11)
hp1.font.bold = True
hp1.font.color.rgb = EMERALD_PRIMARY

hp2 = tf_h.add_paragraph()
hp2.text = "Modern & Accessible Design"
hp2.font.size = Pt(18)
hp2.font.bold = True
hp2.font.color.rgb = SLATE_DARK
hp2.space_before = Pt(4)

hp3 = tf_h.add_paragraph()
hp3.text = "• Hero Banner: Quick action buttons for uploading items or viewing SDG impact metrics."
hp3.font.size = Pt(12)
hp3.font.color.rgb = SLATE_DARK
hp3.space_before = Pt(12)

hp4 = tf_h.add_paragraph()
hp4.text = "• Fast Navigation: Direct access to Analyzer, 100-Item History, Dashboard, and Educational resources."
hp4.font.size = Pt(12)
hp4.font.color.rgb = SLATE_DARK
hp4.space_before = Pt(10)

hp5 = tf_h.add_paragraph()
hp5.text = "• Responsive Layout: Built with Tailwind CSS for seamless mobile and desktop accessibility."
hp5.font.size = Pt(12)
hp5.font.color.rgb = SLATE_DARK
hp5.space_before = Pt(10)

# Right Column: Screenshot
if os.path.exists(HOME_IMG):
    slide4.shapes.add_picture(HOME_IMG, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.5))


# ====================================================
# SLIDE 5: AI Item Analyzer & Classifier (With Screenshot)
# ====================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5, LIGHT_BG)
add_header(slide5, "Artificial Intelligence Core", "AI Item Analyzer & Category Correction Engine")

# Left Column: Features List
create_card(slide5, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), bg_color=CARD_BG)
a_box = slide5.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
tf_a = a_box.text_frame
tf_a.word_wrap = True

ap1 = tf_a.paragraphs[0]
ap1.text = "AI CLASSIFICATION PIPELINE"
ap1.font.size = Pt(11)
ap1.font.bold = True
ap1.font.color.rgb = EMERALD_PRIMARY

ap2 = tf_a.add_paragraph()
ap2.text = "Dual Heuristic & Vision Model"
ap2.font.size = Pt(18)
ap2.font.bold = True
ap2.font.color.rgb = SLATE_DARK
ap2.space_before = Pt(4)

ap3 = tf_a.add_paragraph()
ap3.text = "• MobileNetV2 Vision Model: Analyzes uploaded image features with PyTorch inference."
ap3.font.size = Pt(12)
ap3.font.color.rgb = SLATE_DARK
ap3.space_before = Pt(12)

ap4 = tf_a.add_paragraph()
ap4.text = "• Category Normalization: Maps inputs cleanly into 8 standard target categories."
ap4.font.size = Pt(12)
ap4.font.color.rgb = SLATE_DARK
ap4.space_before = Pt(10)

ap5 = tf_a.add_paragraph()
ap5.text = "• Category Correction Modal: Allows user to override AI detection with instant state sync."
ap5.font.size = Pt(12)
ap5.font.color.rgb = SLATE_DARK
ap5.space_before = Pt(10)

# Right Column: Screenshot
if os.path.exists(ANALYZE_IMG):
    slide5.shapes.add_picture(ANALYZE_IMG, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.5))


# ====================================================
# SLIDE 6: Item Provenance & Printable QR Tag Feature
# ====================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6, LIGHT_BG)
add_header(slide6, "Asset Tracking & Circular Trust", "Printable QR Tag & Provenance Certificate")

# 3 Horizontal Cards Layout
# Card 1: QR Code Generation
create_card(slide6, Inches(0.8), Inches(1.6), Inches(3.644), Inches(5.2), bg_color=CARD_BG)
q1_box = slide6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(3.244), Inches(4.8))
tf_q1 = q1_box.text_frame
tf_q1.word_wrap = True

qp1 = tf_q1.paragraphs[0]
qp1.text = "01. QR TAG GENERATION"
qp1.font.size = Pt(11)
qp1.font.bold = True
qp1.font.color.rgb = EMERALD_PRIMARY

qp2 = tf_q1.add_paragraph()
qp2.text = "High-Res QR Server API"
qp2.font.size = Pt(16)
qp2.font.bold = True
qp2.font.color.rgb = SLATE_DARK
qp2.space_before = Pt(4)

qp3 = tf_q1.add_paragraph()
qp3.text = "• Generates scannable QR tags encoding item ID, title, and provenance link.\n\n• Designed for physical printout on standard label printers to stick onto recycled/donated items."
qp3.font.size = Pt(12)
qp3.font.color.rgb = SLATE_MUTED
qp3.space_before = Pt(12)

# Card 2: Provenance Certificate
create_card(slide6, Inches(4.844), Inches(1.6), Inches(3.644), Inches(5.2), bg_color=CARD_BG)
q2_box = slide6.shapes.add_textbox(Inches(5.044), Inches(1.8), Inches(3.244), Inches(4.8))
tf_q2 = q2_box.text_frame
tf_q2.word_wrap = True

qp4 = tf_q2.paragraphs[0]
qp4.text = "02. DIGITAL PROVENANCE"
qp4.font.size = Pt(11)
qp4.font.bold = True
qp4.font.color.rgb = EMERALD_PRIMARY

qp5 = tf_q2.add_paragraph()
qp5.text = "Item Passport & History"
qp5.font.size = Pt(16)
qp5.font.bold = True
qp5.font.color.rgb = SLATE_DARK
qp5.space_before = Pt(4)

qp6 = tf_q2.add_paragraph()
qp6.text = "• Verifies item authenticity, material category, and AI circularity score.\n\n• Builds trust between donors, charities, and second-hand buyers."
qp6.font.size = Pt(12)
qp6.font.color.rgb = SLATE_MUTED
qp6.space_before = Pt(12)

# Card 3: Community Sharing
create_card(slide6, Inches(8.888), Inches(1.6), Inches(3.644), Inches(5.2), bg_color=CARD_BG)
q3_box = slide6.shapes.add_textbox(Inches(9.088), Inches(1.8), Inches(3.244), Inches(4.8))
tf_q3 = q3_box.text_frame
tf_q3.word_wrap = True

qp7 = tf_q3.paragraphs[0]
qp7.text = "03. CIRCULAR HANDOFF"
qp7.font.size = Pt(11)
qp7.font.bold = True
qp7.font.color.rgb = EMERALD_PRIMARY

qp8 = tf_q3.add_paragraph()
qp8.text = "Print & Tag Utility"
qp8.font.size = Pt(16)
qp8.font.bold = True
qp8.font.color.rgb = SLATE_DARK
qp8.space_before = Pt(4)

qp9 = tf_q3.add_paragraph()
qp9.text = "• Integrated Modal: Features instant 'Print Tag & Certificate' dialog window across Analyze, History, and Item Detail pages."
qp9.font.size = Pt(12)
qp9.font.color.rgb = SLATE_MUTED
qp9.space_before = Pt(12)


# ====================================================
# SLIDE 7: History & SQLite Database (With Screenshot)
# ====================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7, LIGHT_BG)
add_header(slide7, "Data Architecture & Storage", "100-Item Analyzed History & SQLite Backend")

# Left Column: Features List
create_card(slide7, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), bg_color=CARD_BG)
db_box = slide7.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
tf_db = db_box.text_frame
tf_db.word_wrap = True

dbp1 = tf_db.paragraphs[0]
dbp1.text = "DATABASE ARCHITECTURE"
dbp1.font.size = Pt(11)
dbp1.font.bold = True
dbp1.font.color.rgb = EMERALD_PRIMARY

dbp2 = tf_db.add_paragraph()
dbp2.text = "Persistent SQLite Database"
dbp2.font.size = Pt(18)
dbp2.font.bold = True
dbp2.font.color.rgb = SLATE_DARK
dbp2.space_before = Pt(4)

dbp3 = tf_db.add_paragraph()
dbp3.text = "• Database File: SQLite 3 engine located at backend/secondlife.db."
dbp3.font.size = Pt(12)
dbp3.font.color.rgb = SLATE_DARK
dbp3.space_before = Pt(12)

dbp4 = tf_db.add_paragraph()
dbp4.text = "• 100 Seeded Item Records: Complete with realistic condition balance (Damaged: 28, Fair: 28, Good: 29, Not usable: 11, Excellent: 4)."
dbp4.font.size = Pt(12)
dbp4.font.color.rgb = SLATE_DARK
dbp4.space_before = Pt(10)

dbp5 = tf_db.add_paragraph()
dbp5.text = "• Interactive Filtering: Search by keyword, filter by category, or sort by circularity score."
dbp5.font.size = Pt(12)
dbp5.font.color.rgb = SLATE_DARK
dbp5.space_before = Pt(10)

# Right Column: Screenshot
if os.path.exists(HISTORY_IMG):
    slide7.shapes.add_picture(HISTORY_IMG, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.5))


# ====================================================
# SLIDE 8: Sustainability Impact Dashboard (With Screenshot)
# ====================================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8, LIGHT_BG)
add_header(slide8, "Analytics & Metrics", "Sustainability Dashboard & Circularity Scoring")

# Left Column: Features List
create_card(slide8, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), bg_color=CARD_BG)
d_box = slide8.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
tf_d = d_box.text_frame
tf_d.word_wrap = True

dp1 = tf_d.paragraphs[0]
dp1.text = "ANALYTICS & METRICS"
dp1.font.size = Pt(11)
dp1.font.bold = True
dp1.font.color.rgb = EMERALD_PRIMARY

dp2 = tf_d.add_paragraph()
dp2.text = "Circularity Scoring (0-100)"
dp2.font.size = Pt(18)
dp2.font.bold = True
dp2.font.color.rgb = SLATE_DARK
dp2.space_before = Pt(4)

dp3 = tf_d.add_paragraph()
dp3.text = "• Circularity Index: Evaluates item reusability, material recyclability, and repair potential."
dp3.font.size = Pt(12)
dp3.font.color.rgb = SLATE_DARK
dp3.space_before = Pt(12)

dp4 = tf_d.add_paragraph()
dp4.text = "• Visual KPI Cards: Real-time tracking of total items analyzed, waste diverted, and donations."
dp4.font.size = Pt(12)
dp4.font.color.rgb = SLATE_DARK
dp4.space_before = Pt(10)

dp5 = tf_d.add_paragraph()
dp5.text = "• Category Distribution: Interactive visual breakdown across all 8 target categories."
dp5.font.size = Pt(12)
dp5.font.color.rgb = SLATE_DARK
dp5.space_before = Pt(10)

# Right Column: Screenshot
if os.path.exists(DASHBOARD_IMG):
    slide8.shapes.add_picture(DASHBOARD_IMG, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.5))


# ====================================================
# SLIDE 9: Detailed SDG Framework Mapping (With Screenshot)
# ====================================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9, LIGHT_BG)
add_header(slide9, "Academic Integrity & Framework", "UN Sustainable Development Goals Mapping Page")

# Left Column: Features List
create_card(slide9, Inches(0.8), Inches(1.4), Inches(4.5), Inches(5.5), bg_color=CARD_BG)
sd_box = slide9.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(4.1), Inches(5.1))
tf_sd = sd_box.text_frame
tf_sd.word_wrap = True

sdp1 = tf_sd.paragraphs[0]
sdp1.text = "SDG IMPACT PAGE"
sdp1.font.size = Pt(11)
sdp1.font.bold = True
sdp1.font.color.rgb = EMERALD_PRIMARY

sdp2 = tf_sd.add_paragraph()
sdp2.text = "Transparent Framework"
sdp2.font.size = Pt(18)
sdp2.font.bold = True
sdp2.font.color.rgb = SLATE_DARK
sdp2.space_before = Pt(4)

sdp3 = tf_sd.add_paragraph()
sdp3.text = "• Dedicated SDG Page: Full breakdown of SDG 12, SDG 11, and SDG 13 targets."
sdp3.font.size = Pt(12)
sdp3.font.color.rgb = SLATE_DARK
sdp3.space_before = Pt(12)

sdp4 = tf_sd.add_paragraph()
sdp4.text = "• Material Energy Savings: Highlights facts like 95% energy saved by recycling aluminum."
sdp4.font.size = Pt(12)
sdp4.font.color.rgb = SLATE_DARK
sdp4.space_before = Pt(10)

sdp5 = tf_sd.add_paragraph()
sdp5.text = "• Academic Integrity: Explains qualitative impact scoring without claiming unverified carbon metrics."
sdp5.font.size = Pt(12)
sdp5.font.color.rgb = SLATE_DARK
sdp5.space_before = Pt(10)

# Right Column: Screenshot
if os.path.exists(SDG_IMG):
    slide9.shapes.add_picture(SDG_IMG, Inches(5.6), Inches(1.4), Inches(6.933), Inches(5.5))


# ====================================================
# SLIDE 10: Technical Architecture & Future Roadmap
# ====================================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10, DARK_BG)

# Title Header
add_header(slide10, "Technology & Horizon", "Technical Architecture & Future Roadmap")

# Left Column: Tech Stack Card
create_card(slide10, Inches(0.8), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=RGBColor(30, 41, 59), border_color=EMERALD_PRIMARY)
t_box = slide10.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.3), Inches(5.1))
tf_tech = t_box.text_frame
tf_tech.word_wrap = True

tp1 = tf_tech.paragraphs[0]
tp1.text = "TECHNOLOGY STACK"
tp1.font.size = Pt(12)
tp1.font.bold = True
tp1.font.color.rgb = EMERALD_PRIMARY

tp2 = tf_tech.add_paragraph()
tp2.text = "• Frontend: React 18, Vite, Tailwind CSS, Lucide Icons, React Router."
tp2.font.size = Pt(13)
tp2.font.color.rgb = RGBColor(255, 255, 255)
tp2.space_before = Pt(12)

tp3 = tf_tech.add_paragraph()
tp3.text = "• Backend API: FastAPI (Python 3), Uvicorn ASGI Server, CORS Middleware."
tp3.font.size = Pt(13)
tp3.font.color.rgb = RGBColor(255, 255, 255)
tp3.space_before = Pt(10)

tp4 = tf_tech.add_paragraph()
tp4.text = "• AI Engine: PyTorch MobileNetV2 Vision Model + Word-Boundary Keyword Normalization."
tp4.font.size = Pt(13)
tp4.font.color.rgb = RGBColor(255, 255, 255)
tp4.space_before = Pt(10)

tp5 = tf_tech.add_paragraph()
tp5.text = "• Database & Services: SQLite 3 Engine (`secondlife.db`), QR Server API Integration."
tp5.font.size = Pt(13)
tp5.font.color.rgb = RGBColor(255, 255, 255)
tp5.space_before = Pt(10)

# Right Column: Future Roadmap Card
create_card(slide10, Inches(6.833), Inches(1.4), Inches(5.7), Inches(5.5), bg_color=RGBColor(30, 41, 59), border_color=AMBER_PRIMARY)
r_box = slide10.shapes.add_textbox(Inches(7.033), Inches(1.6), Inches(5.3), Inches(5.1))
tf_road = r_box.text_frame
tf_road.word_wrap = True

rp1 = tf_road.paragraphs[0]
rp1.text = "FUTURE ROADMAP"
rp1.font.size = Pt(12)
rp1.font.bold = True
rp1.font.color.rgb = AMBER_PRIMARY

rp2 = tf_road.add_paragraph()
rp2.text = "1. Geolocation Recycling Map: Integrate OpenStreetMap / Leaflet to display nearest drop-off centers."
rp2.font.size = Pt(13)
rp2.font.color.rgb = RGBColor(255, 255, 255)
rp2.space_before = Pt(12)

rp3 = tf_road.add_paragraph()
rp3.text = "2. Community Marketplace API: Direct P2P item donation and second-hand exchange listing."
rp3.font.size = Pt(13)
rp3.font.color.rgb = RGBColor(255, 255, 255)
rp3.space_before = Pt(10)

rp4 = tf_road.add_paragraph()
rp4.text = "3. IoT Smart Sorting Bins: Connect camera modules to physical waste sorting bins for automatic item identification."
rp4.font.size = Pt(13)
rp4.font.color.rgb = RGBColor(255, 255, 255)
rp4.space_before = Pt(10)


# ----------------------------------------------------
# Save Presentation Output
# ----------------------------------------------------
output_path = r"c:\Users\MINHA FATHIMA C\OneDrive\Desktop\ai\SecondLife_AI_Presentation.pptx"
prs.save(output_path)
print(f"Presentation successfully saved to: {output_path}")
